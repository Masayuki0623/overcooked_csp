"""進捗報告のスライドを作る。"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pathlib import Path

OUT = Path(r'C:/Users/sanda/PythonPrograms/overcooked_csp/docs/progress_report.pptx')

TEAL = RGBColor(0x02, 0x80, 0x90)
SEA = RGBColor(0x00, 0xA8, 0x96)
MINT = RGBColor(0x02, 0xC3, 0x9A)
DARK = RGBColor(0x05, 0x32, 0x3A)
INK = RGBColor(0x1A, 0x2E, 0x33)
GREY = RGBColor(0x6B, 0x7C, 0x80)
PALE = RGBColor(0xEC, 0xF4, 0xF4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CORAL = RGBColor(0xD9, 0x53, 0x4F)

FONT = 'Meiryo'
W, H = 13.333, 7.5

prs = Presentation()
prs.slide_width = Inches(W)  # 16:9
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]


def slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    bgshape = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                 Inches(W), Inches(H))
    bgshape.fill.solid()
    bgshape.fill.fore_color.rgb = bg
    bgshape.line.fill.background()
    bgshape.shadow.inherit = False
    return s


def text(s, x, y, w, h, runs, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT,
         space_after=6, line=None, anchor=MSO_ANCHOR.TOP):
    """runs: 文字列、または (文字列, dict) のリスト。"""
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    items = runs if isinstance(runs, list) else [runs]
    for i, item in enumerate(items):
        body, opt = item if isinstance(item, tuple) else (item, {})
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opt.get('align', align)
        p.space_after = Pt(opt.get('space_after', space_after))
        if line or opt.get('line'):
            p.line_spacing = opt.get('line', line)
        r = p.add_run()
        r.text = body
        f = r.font
        f.name = FONT
        f.size = Pt(opt.get('size', size))
        f.bold = opt.get('bold', bold)
        f.color.rgb = opt.get('color', color)
    return box


def card(s, x, y, w, h, fill=PALE, radius=True, line_color=None):
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line_color is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line_color
        shp.line.width = Pt(1)
    shp.shadow.inherit = False
    if radius:
        try:
            shp.adjustments[0] = 0.08
        except Exception:
            pass
    shp.text_frame.word_wrap = True
    return shp


def badge(s, x, y, d, label, fill=TEAL, fg=WHITE, size=16):
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = fill
    c.line.fill.background()
    c.shadow.inherit = False
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = fg
    return c


def title(s, t, sub=None, color=INK, subcolor=GREY):
    # 全角で26字を超えると32ptでは1行に収まらず、副題に重なる。
    size = 32 if len(t) <= 25 else 27
    text(s, 0.75, 0.55, 11.8, 0.85, t, size=size, bold=True, color=color)
    if sub:
        text(s, 0.75, 1.42, 11.8, 0.45, sub, size=15, color=subcolor)


# ---------------------------------------------------------------- 1 表紙
s = slide(DARK)
text(s, 1.0, 2.25, 11.3, 0.5, '研究進捗報告', size=17, color=MINT, bold=True)
text(s, 1.0, 2.95, 11.3, 1.9,
     ['AI はどこまで人の指示を',
      '「後回し」にしてよいのか'],
     size=42, bold=True, color=WHITE, line=1.25, space_after=0)
text(s, 1.0, 5.15, 11.3, 0.9,
     ['人と AI が協力して料理をつくる環境で、',
      'AI が指示を先送りする度合いが、人の体験にどう響くかを調べています。'],
     size=15, color=RGBColor(0xB9, 0xD6, 0xD9), line=1.35, space_after=2)
text(s, 1.0, 6.55, 11.3, 0.4, '2026年9月9日', size=12, color=GREY)

# ---------------------------------------------------------------- 2 一言でいうと
s = slide()
title(s, 'この研究を一言でいうと', '効率のよい AI ほど、人の指示を後回しにする。その代償を測ります。')

items = [
    ('1', '状況', '人と AI が2人で協力し、3つの料理を最短時間で仕上げる。'
                  'AI は全体最適の計画を立てて動く。'),
    ('2', '起きること', '人が「これをやって」と頼んでも、AI の計画に合わない指示は'
                        '後回しにされる。効率を優先するほどそうなる。'),
    ('3', '問い', '後回しにされたとき、人はどれだけ「聞いてもらえなかった」と'
                  '感じるのか。その度合いを操作して測る。'),
]
for i, (num, head, body) in enumerate(items):
    y = 2.15 + i * 1.62
    card(s, 0.75, y, 11.8, 1.35)
    badge(s, 1.05, y + 0.36, 0.62, num)
    text(s, 1.95, y + 0.26, 2.0, 0.4, head, size=17, bold=True, color=TEAL)
    text(s, 4.05, y + 0.28, 8.2, 0.9, body, size=14, color=INK, line=1.3)

# ---------------------------------------------------------------- 3 題材
s = slide()
title(s, '題材：仕切りで分かれたキッチン', '2人は行き来できない。相手に頼まないと完成しない料理がある。')

# 左パネル(AI)
card(s, 0.75, 2.1, 4.55, 3.9, PALE)
text(s, 1.05, 2.32, 3.95, 0.4, 'AI 側（左）', size=16, bold=True, color=TEAL)
text(s, 1.05, 2.9, 3.95, 2.9,
     ['具材　レタス・玉ねぎ・りんご・オレンジ',
      '設備　鍋',
      'まな板　なし'],
     size=14, color=INK, line=1.5, space_after=10)
text(s, 1.05, 5.15, 3.95, 0.7,
     'AI 側の具材は、そのままでは切れない。',
     size=13, color=CORAL, bold=True, line=1.3)

# 中央(共有台)
mid = card(s, 5.55, 2.1, 2.2, 3.9, RGBColor(0xD7, 0xE9, 0xEA))
text(s, 5.62, 3.35, 2.06, 1.4,
     ['共有台', '（仕切りの上）', '両側から使える'],
     size=13, bold=True, color=TEAL, align=PP_ALIGN.CENTER, line=1.4, space_after=4)

# 右パネル(人間)
card(s, 8.0, 2.1, 4.55, 3.9, PALE)
text(s, 8.3, 2.32, 3.95, 0.4, '人間 側（右）', size=16, bold=True, color=TEAL)
text(s, 8.3, 2.9, 3.95, 2.9,
     ['具材　トマト・バナナ',
      '設備　ミキサー・提供口',
      'まな板　あり（1つだけ）'],
     size=14, color=INK, line=1.5, space_after=10)
text(s, 8.3, 5.15, 3.95, 0.7,
     '切る作業は、すべてここでしかできない。',
     size=13, color=CORAL, bold=True, line=1.3)

text(s, 0.75, 6.25, 11.8, 0.75,
     'だから AI は「材料を中央まで運ぶ」ことしかできず、'
     '人間が切らないと料理が進まない。互いに依存する状況が生まれる。',
     size=14, color=INK, line=1.35)

# ---------------------------------------------------------------- 4 対立の構造
s = slide()
title(s, 'なぜ指示が「合う / 合わない」に分かれるのか',
      '3品のうちスープだけが全体の所要時間を決める。そこに全体最適と個人の都合の対立がある。')

card(s, 0.75, 2.2, 5.75, 3.5, PALE)
badge(s, 1.1, 2.5, 0.6, '◎', TEAL)
text(s, 1.95, 2.62, 4.3, 0.45, 'スープを進める指示', size=18, bold=True, color=TEAL)
text(s, 1.1, 3.4, 5.05, 2.1,
     ['煮込みに15秒かかり、全体の所要時間を決めている。',
      '早く着手しないと全体が遅れる。',
      'AI の計画と一致するため、すぐ実行される。'],
     size=14, color=INK, line=1.4, space_after=9)

card(s, 6.85, 2.2, 5.7, 3.5, PALE)
badge(s, 7.2, 2.5, 0.6, '△', CORAL)
text(s, 8.05, 2.62, 4.3, 0.45, 'ジュースを進める指示', size=18, bold=True, color=CORAL)
text(s, 7.2, 3.4, 5.0, 2.1,
     ['人間側で完結しやすく、急ぐ必要がない。',
      '人にとっては「自分の作業を早く始めたい」',
      'という自然な要求。',
      'AI の計画とずれるため、後回しにされる。'],
     size=14, color=INK, line=1.4, space_after=9)

text(s, 0.75, 6.1, 11.8, 0.9,
     ['参加者にとっては、どちらも等しくまっとうな要求です。'
      '差が生まれるのは AI の計画にとってであり、それは参加者には見えません。'],
     size=14, color=INK, bold=True, line=1.35)

# ---------------------------------------------------------------- 5 独立変数
s = slide()
title(s, '操作する条件：AI が指示を後回しにできる量',
      '「指示された作業に取りかかる前に、他の作業をいくつ挟んでよいか」を 0 / 2 / 4 で変える。')

levels = [
    ('0', 'すぐやる', '他の作業を挟めない。\n指示は必ず最優先で実行される。', TEAL),
    ('2', 'すこし後回し', '他の作業を2つまで挟める。\n指示は3番目あたりに落ちる。', SEA),
    ('4', 'かなり後回し', '他の作業を4つまで挟める。\n指示しない場合とほぼ同じ順序になる。', CORAL),
]
for i, (num, head, body, col) in enumerate(levels):
    x = 0.75 + i * 4.05
    card(s, x, 2.3, 3.75, 3.2, PALE)
    badge(s, x + 0.3, 2.6, 0.85, num, col, WHITE, 26)
    text(s, x + 1.35, 2.82, 2.2, 0.45, head, size=16, bold=True, color=col)
    text(s, x + 0.3, 3.85, 3.15, 1.4, body.replace('\n', ''),
         size=13.5, color=INK, line=1.45)

text(s, 0.75, 5.95, 11.8, 0.9,
     '秒数ではなく「作業の個数」で決めます。移動距離や作業の重さで体感時間がばらつくのを避けるためです。'
     'また、その指示に必要な前準備は数えません。',
     size=14, color=INK, line=1.35)

# ---------------------------------------------------------------- 6 測るもの
s = slide()
title(s, '何を測るか', '主役は「体験」。効率は、対抗する説明を潰すために測ります。')

card(s, 0.75, 2.2, 5.75, 3.9, RGBColor(0xE3, 0xF2, 0xF1))
badge(s, 1.1, 2.5, 0.62, '主', TEAL)
text(s, 1.98, 2.63, 4.2, 0.45, '体験（アンケート）', size=18, bold=True, color=TEAL)
text(s, 1.1, 3.45, 5.05, 2.4,
     ['協調感　CCR スケール（6項目）',
      '信頼感　MDMT スケール（8項目）',
      '',
      '各セッション終了直後に回答。所要1分。'],
     size=14, color=INK, line=1.45, space_after=7)

card(s, 6.85, 2.2, 5.7, 3.9, PALE)
badge(s, 7.2, 2.5, 0.62, '客', GREY)
text(s, 8.08, 2.63, 4.2, 0.45, '客観データ（自動記録）', size=18, bold=True, color=INK)
text(s, 7.2, 3.45, 5.0, 2.4,
     ['指示してから着手までの待ち時間',
      '指示された作業が何番目に実行されたか',
      '待っている間に AI がこなした他の作業',
      '3品完成までの所要時間'],
     size=14, color=INK, line=1.45, space_after=7)

text(s, 0.75, 6.35, 11.8, 0.6,
     '仮説：後回しの量が小さすぎても大きすぎても体験は悪くなる（逆U字）。',
     size=15, bold=True, color=TEAL)

# ---------------------------------------------------------------- 7 これまでの作業
s = slide()
title(s, 'ここまでにやったこと：環境を実験に使える状態にする',
      '人に見せる前に、シミュレーションで環境が壊れていないかを徹底的に確かめました。')

stats = [
    ('17件', '実装の不具合を修正', '大半が同じ系統。「世界に問い合わせれば分かること」を、\n名前や配列の位置から推測していた箇所。'),
    ('37% → 100%', '料理を完成できた割合', '当初は6割の試行が時間切れ。時間切れは条件に偏っており、\nそのままでは条件間の比較が成立しなかった。'),
    ('39件 → 0件', '2人とも動かない時間', '完走はするが両者が最大23秒立ち尽くす試行があった。\n完走率だけを見ていては気づけなかった。'),
]
for i, (num, head, body) in enumerate(stats):
    y = 2.2 + i * 1.53
    card(s, 0.75, y, 11.8, 1.3)
    text(s, 1.05, y + 0.3, 2.75, 0.7, num, size=25, bold=True, color=TEAL)
    text(s, 4.0, y + 0.18, 2.9, 0.4, head, size=14, bold=True, color=INK)
    text(s, 4.0, y + 0.62, 8.3, 0.6, body.replace('\n', ''),
         size=12.5, color=GREY, line=1.3)

text(s, 0.75, 6.9, 11.8, 0.45,
     '不具合は、左右を仕切ったキッチンに変えたときに一斉に表面化しました。',
     size=13, color=GREY)

# ---------------------------------------------------------------- 8 健全性
s = slide(DARK)
text(s, 0.75, 0.7, 11.8, 0.7, '現在の環境の状態', size=32, bold=True, color=WHITE)
text(s, 0.75, 1.55, 11.8, 0.5,
     '本実験で使う12通りの注文構成で、216通りの条件をすべて確認しました。',
     size=15, color=RGBColor(0xB9, 0xD6, 0xD9))

facts = [('216 / 216', '3品すべて完成'), ('0 件', '2人とも止まる時間'),
         ('29.6 秒', '平均の所要時間'), ('41.4 秒', '最も遅い試行')]
for i, (num, lab) in enumerate(facts):
    x = 0.75 + i * 3.02
    c = card(s, x, 2.6, 2.75, 1.95, RGBColor(0x0A, 0x46, 0x50))
    text(s, x + 0.2, 2.95, 2.35, 0.75, num, size=26, bold=True, color=MINT,
         align=PP_ALIGN.CENTER)
    text(s, x + 0.2, 3.82, 2.35, 0.5, lab, size=13,
         color=RGBColor(0xB9, 0xD6, 0xD9), align=PP_ALIGN.CENTER)

text(s, 0.75, 5.0, 11.8, 1.5,
     ['制限時間は100秒。最も遅い試行でも41.4秒なので、余裕があります。',
      '「2人とも止まる」は、両者の位置も持ち物も3秒以上変わらない状態のこと。'
      '煮込みの待ち時間は正当な待機として除外して数えています。'],
     size=14, color=RGBColor(0xD3, 0xE7, 0xE9), line=1.45, space_after=9)

# ---------------------------------------------------------------- 9 発見1（グラフ）
s = slide()
title(s, '分かったこと①　指示は順序をほとんど変えなくなる',
      '後回しできる量が大きいほど。数字は「指示によって何番目ぶん早くなったか」で、'
      '0 なら指示してもしなくても同じ。')

cd = CategoryChartData()
cd.categories = ['0（すぐやる）', '2（すこし後回し）', '4（かなり後回し）']
cd.add_series('計画に合う指示', (1.25, 0.00, 0.00))
cd.add_series('計画に合わない指示', (3.17, 1.17, 0.54))
cd.add_series('無作為な指示', (2.08, 0.50, -0.23))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.75), Inches(2.15),
                        Inches(7.6), Inches(4.0), cd)
ch = gf.chart
ch.has_title = False
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.BOTTOM
ch.legend.include_in_layout = False
ch.legend.font.size = Pt(11)
ch.legend.font.name = FONT
for srs, col in zip(ch.plots[0].series, (TEAL, CORAL, GREY)):
    srs.format.fill.solid()
    srs.format.fill.fore_color.rgb = col
pl = ch.plots[0]
pl.has_data_labels = True
pl.data_labels.number_format = '0.00'
pl.data_labels.number_format_is_linked = False
pl.data_labels.font.size = Pt(10)
pl.data_labels.font.name = FONT
pl.data_labels.font.color.rgb = INK
for ax in (ch.category_axis, ch.value_axis):
    ax.tick_labels.font.size = Pt(11)
    ax.tick_labels.font.name = FONT
    ax.tick_labels.font.color.rgb = GREY
ch.value_axis.has_major_gridlines = True
ch.category_axis.has_major_gridlines = False

card(s, 8.7, 2.15, 3.85, 4.0, PALE)
text(s, 9.05, 2.5, 3.15, 3.3,
     [('後回しできる量が 4 のとき', {'size': 15, 'bold': True, 'color': TEAL}),
      ('計画に合わない指示は +0.54。指示しなかった場合と'
       'ほぼ同じ順序になります。', {'size': 13.5}),
      ('', {'size': 6}),
      ('参加者は「聞いてはいるが、自分の意図は反映されていない」'
       '状態に置かれます。', {'size': 13.5, 'bold': True, 'color': INK})],
     line=1.45, space_after=11)
text(s, 8.7, 6.35, 3.85, 0.7,
     'この研究が想定する心理的な仕組みに対応する、客観的な事実です。',
     size=12.5, color=GREY, line=1.35)

# ---------------------------------------------------------------- 10 発見2
s = slide()
title(s, '分かったこと②　後回しにしても、料理は遅くならない',
      '「遅くなるから不満なのだ」という別の説明を、あらかじめ潰しておくための確認です。')

rows = [
    ('相手の動き方', '0', '2', '4', '0→4 の変化'),
    ('人間らしい相手', '35.4 秒', '31.8 秒', '34.4 秒', '−0.96 秒（差なし）'),
    ('理想的な相手', '25.6 秒', '25.5 秒', '25.1 秒', '−0.44 秒（差なし）'),
]
colw = [3.0, 1.85, 1.85, 1.85, 3.25]
x0, y0 = 0.75, 2.3
for ri, row in enumerate(rows):
    x = x0
    for ci, cell in enumerate(row):
        fill = TEAL if ri == 0 else (PALE if ri % 2 else RGBColor(0xF7, 0xFB, 0xFB))
        c = card(s, x, y0 + ri * 0.72, colw[ci] - 0.06, 0.66, fill, radius=False)
        text(s, x + 0.15, y0 + ri * 0.72 + 0.19, colw[ci] - 0.36, 0.4, cell,
             size=13.5, bold=(ri == 0),
             color=WHITE if ri == 0 else INK,
             align=PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER)
        x += colw[ci]

card(s, 0.75, 4.75, 11.8, 1.35, RGBColor(0xE3, 0xF2, 0xF1))
text(s, 1.1, 5.0, 11.1, 0.9,
     [('どちらの相手でも、統計的に意味のある変化はありませんでした。',
       {'size': 15, 'bold': True, 'color': TEAL}),
      ('つまり、もし体験の悪化が観測されれば、それは「料理が遅くなったから」では説明できません。',
       {'size': 14, 'color': INK})],
     line=1.4, space_after=6)

text(s, 0.75, 6.4, 11.8, 0.85,
     'この数値は3度変わりました。停止時間による汚染、あいまいな注文構成の混入を'
     'それぞれ特定して取り除いた結果が上の値です。',
     size=12.5, color=GREY, line=1.35)

# ---------------------------------------------------------------- 11 これから
s = slide()
title(s, '現在地と、これから', '環境の準備は終わりました。ここから人に参加してもらいます。')

done = ['環境の実装と検証（完走率100%・停止0件）',
        'シミュレーションによる事前検証（216通り）',
        'アンケート項目の確定（既存の検証済み尺度）',
        '参加者がプレイする経路の用意']
todo = ['自分で3条件をプレイして、違いが分かるか確認',
        '人の動きとシミュレーションの人間モデルを突き合わせ',
        '参加者の募集（15〜30名）',
        '本実験の実施と分析']

card(s, 0.75, 2.2, 5.75, 4.15, RGBColor(0xE3, 0xF2, 0xF1))
badge(s, 1.1, 2.5, 0.6, '済', TEAL)
text(s, 1.98, 2.63, 4.2, 0.45, '完了していること', size=18, bold=True, color=TEAL)
text(s, 1.1, 3.45, 5.05, 2.6, ['・' + t for t in done],
     size=14, color=INK, line=1.4, space_after=13)

card(s, 6.85, 2.2, 5.7, 4.15, PALE)
badge(s, 7.2, 2.5, 0.6, '次', CORAL)
text(s, 8.08, 2.63, 4.2, 0.45, 'これからやること', size=18, bold=True, color=CORAL)
text(s, 7.2, 3.45, 5.0, 2.6, ['・' + t for t in todo],
     size=14, color=INK, line=1.4, space_after=13)

text(s, 0.75, 6.6, 11.8, 0.5,
     '参加者は1人あたり3セッション（後回しの量 0 / 2 / 4）を体験し、毎回アンケートに答えます。',
     size=13.5, color=GREY)

# ---------------------------------------------------------------- 12 まとめ
s = slide(DARK)
text(s, 1.0, 1.75, 11.3, 0.5, 'まとめ', size=17, bold=True, color=MINT)
text(s, 1.0, 2.5, 11.3, 2.6,
     ['効率のよい AI ほど、人の指示を後回しにします。',
      'その「後回し」が、人の協調感と信頼感にどう響くのか。',
      'それを測れる環境が、ようやく整いました。'],
     size=27, bold=True, color=WHITE, line=1.4, space_after=10)
text(s, 1.0, 5.6, 11.3, 0.9,
     '次は、実際に人に遊んでもらいます。',
     size=16, color=RGBColor(0xB9, 0xD6, 0xD9))

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(str(OUT))
print('保存: %s (%d枚)' % (OUT, len(prs.slides.__iter__.__self__._sldIdLst)))
